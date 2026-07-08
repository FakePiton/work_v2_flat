import os
from docx import Document
from babel.dates import format_date
from datetime import datetime, timedelta
from src.services.data_repository import PandasDataRepository


class NewOrder:
    def __init__(self, pd_data_repository: PandasDataRepository):
        self.pd_data_repository = pd_data_repository
        self.text_info = ""

    @staticmethod
    def format_ukrainian_date(date):
        day = date.strftime("%d")
        month_year = format_date(date, format="MMMM yyyy", locale='uk')
        return f"{day} {month_year}"

    def get_path_server(self, date: datetime, path_dir_server: str):
        dir_name_year = f"{date.year}(samba)"
        dir_name_month = f"{date.strftime('%m')} накази"

        path_year = os.path.join(path_dir_server, dir_name_year)

        if not os.path.exists(path_year):
            self.text_info += f"Папка не найдена. Створюємо папку: {path_year} \n"
            os.makedirs(path_year)

        path_mount = os.path.join(path_dir_server, dir_name_year, dir_name_month)
        if not os.path.exists(path_mount):
            self.text_info += f"Папка не найдена. Створюємо папку: {path_mount} \n"
            os.makedirs(path_mount)

        return path_mount

    def _replace_text_in_doc(self, doc, replacements):
        for paragraph in doc.paragraphs:
            for old, new in replacements.items():
                if old in paragraph.text:
                    paragraph.text = paragraph.text.replace(old, str(new))

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for old, new in replacements.items():
                        if old in cell.text:
                            cell.text = cell.text.replace(old, str(new))

    def create_template(self, path_dir_server: str, path_order_template: str):
        now = datetime.now()
        tomorrow = now + timedelta(days=1)
        formatted_date = self.format_ukrainian_date(tomorrow)

        number = self.pd_data_repository.get_order_number_by_date(date=now.date())
        today_str = now.date().strftime("%d.%m.%Y")

        replacements = {
            "{{ date }}": today_str,
            "{{ number }}": number,
            "{{ date_prod }}": formatted_date,
        }

        path_server = self.get_path_server(date=now, path_dir_server=path_dir_server)

        file_name = f"НАКАЗ №{number} від {today_str} в процесі.docx"
        new_file = os.path.join(path_server, file_name)

        if os.path.isfile(new_file):
            self.text_info += f"📄 Файл вже існує: {new_file}\n"
            return None

        doc = Document(path_order_template)

        self._replace_text_in_doc(doc, replacements)

        doc.save(new_file)

        self.text_info += f"✅ Шаблон створено: {file_name}"
